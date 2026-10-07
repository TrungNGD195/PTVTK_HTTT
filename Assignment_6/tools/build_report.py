"""Build the editable English report from the current project evidence.

Running this script replaces Bao_cao_Assignment_06.docx. Do not rerun after
manually adding screenshots or editing the report unless you have a backup.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
EVALUATION = json.loads((ROOT / "results" / "evaluation.json").read_text(encoding="utf-8"))
DESTINATION = ROOT / "Bao_cao_Assignment_06.docx"


def shade(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    properties.append(element)


def set_cell_border(cell, color: str = "D62828", size: int = 18) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        line = OxmlElement(f"w:{edge}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), str(size))
        line.set(qn("w:color"), color)
        borders.append(line)
    properties.append(borders)


def paragraph(doc: Document, text: str = "", bold_lead: str | None = None):
    item = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        item.add_run(bold_lead).bold = True
        item.add_run(text[len(bold_lead):])
    else:
        item.add_run(text)
    return item


def heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def table(doc: Document, headers: list[str], rows: list[list[str]], widths=None) -> None:
    result = doc.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    result.autofit = True
    for i, label in enumerate(headers):
        cell = result.rows[0].cells[i]
        cell.text = label
        shade(cell, "17365D")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        cells = result.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if len(result.rows) % 2 == 0:
            for cell in cells:
                shade(cell, "EEF3F8")
    for row in result.rows:
        for cell in row.cells:
            for item in cell.paragraphs:
                item.paragraph_format.space_after = Pt(0)
                for run in item.runs:
                    run.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def placeholder(doc: Document, label: str, instructions: str, height_inches: float = 1.25) -> None:
    box = doc.add_table(rows=1, cols=1)
    box.autofit = False
    cell = box.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(cell, "FFF2A8")
    set_cell_border(cell)
    box.rows[0].height = Inches(height_inches)
    item = cell.paragraphs[0]
    item.alignment = WD_ALIGN_PARAGRAPH.CENTER
    top = item.add_run("[ INSERT IMAGE HERE ]\n" + label)
    top.bold = True
    top.font.size = Pt(12)
    top.font.color.rgb = RGBColor(180, 0, 0)
    help_text = cell.add_paragraph(instructions)
    help_text.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in help_text.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(90, 45, 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def new_page(doc: Document, title: str) -> None:
    doc.add_page_break()
    heading(doc, title)


def add_page_number(section) -> None:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Assignment 06  •  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    add_page_number(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.3)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color in (("Title", 24, "17365D"), ("Heading 1", 15, "17365D"),
                              ("Heading 2", 11.5, "1D5385")):
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)
    doc.core_properties.title = "Assignment 06 — Multimodal Search System for E-Commerce"
    doc.core_properties.subject = "Information System Analysis and Design"
    doc.core_properties.author = "Student"

    # Page 1: cover and abstract.
    doc.add_paragraph("INFORMATION SYSTEM ANALYSIS AND DESIGN", style="Subtitle")
    doc.add_heading("Assignment 06", 0)
    doc.add_paragraph("Multimodal Search System for E-Commerce", style="Title")
    paragraph(doc, "Editable project report")
    table(doc, ["Submission detail", "Fill in before submission"], [
        ["Student name / ID", "[ENTER NAME AND STUDENT ID]"],
        ["Class / lecturer", "[ENTER CLASS AND LECTURER]"],
        ["Submission version / date", "[ENTER VERSION AND DATE]"],
    ])
    heading(doc, "Abstract", 2)
    paragraph(doc, "This report presents a three-layer prototype for e-commerce product search using text, simulated voice, and image input. Each mode is normalized into a common query object. Candidate retrieval and result ranking are implemented as separate operations. The prototype searches a 12-product JSON dataset and a small in-memory image vector index; it also demonstrates a combined text-and-image query and customer-scoped order lookup. A 13-query Top-1 evaluation produced 12 successful results (92.31%). The remaining failure illustrates the limits of keyword matching when the customer uses synonyms. The figures and program screenshots marked in yellow are reserved for student-supplied Visual Paradigm screenshots and execution captures.")
    heading(doc, "Keywords", 2)
    paragraph(doc, "multimodal search; three-layer architecture; UML; keyword retrieval; image similarity; ranking")

    # Page 2.
    new_page(doc, "1. Introduction and Problem Description")
    heading(doc, "1.1 Context", 2)
    paragraph(doc, "Online customers may express the same shopping need in different ways: typed keywords, a spoken sentence, or a picture of a desired product. A search system must convert these inputs into comparable representations, find candidate items, and present the most useful results. The assignment focuses on the system design and a working small prototype, rather than a production e-commerce platform.")
    heading(doc, "1.2 Problem and objective", 2)
    paragraph(doc, "The problem is to support several input modes without creating a separate end-to-end architecture for each mode. The proposed solution uses one SearchUI, modality-specific conversion services, a common query dictionary, retrieval services, and a distinct ranking service. Products and orders are stored behind repositories. The user receives item names, prices, and ranking scores, or an order result when a known order ID is supplied.")
    heading(doc, "1.3 Scope", 2)
    bullet(doc, "Required modes: English keyword search, voice search simulated from transcript text, and image-similarity search from a local PNG/JPEG file.")
    bullet(doc, "Additional demonstrations: a weighted text-and-image query, a basic price ceiling, and order lookup constrained by customer ID.")
    bullet(doc, "Dataset: 12 products, 12 product illustrations, two separate query illustrations, and three sample orders.")
    bullet(doc, "Interface: a command-line SearchUI. No web server, microphone capture, or persistent database is required for this prototype.")
    heading(doc, "1.4 Design principle", 2)
    paragraph(doc, "The processing path is: input → common query representation → candidate retrieval → ranking → results. The presentation layer collects input and displays output; the application layer interprets and searches; the data layer stores products, orders, and vectors. This layering prevents direct UI-to-database access and makes each search method replaceable.")

    # Page 3.
    new_page(doc, "2. Requirements Analysis")
    heading(doc, "2.1 Actor, inputs, and outputs", 2)
    paragraph(doc, "The primary actor is the Customer. Inputs are (1) English keyword or natural-language text, (2) transcript text representing a voice request, (3) a local product image, and optionally (4) an order ID plus customer ID. Product-search output consists of ranked product records with product ID, name, price, and score. Order-search output consists of a matching order ID, status, total, and items, or an explicit not-found response.")
    heading(doc, "2.2 Functional requirements", 2)
    table(doc, ["ID", "Requirement", "Implementation"], [
        ["FR1", "Search products by typed text", "Text query + keyword retrieval"],
        ["FR2", "Accept simulated voice input", "SpeechService transcript + voice query"],
        ["FR3", "Search products by image", "ImageService + cosine vector index"],
        ["FR4", "Use a common query format", "QueryService dictionaries with type field"],
        ["FR5", "Rank and show product results", "RankingService + SearchUI scores"],
        ["FR6", "View product information", "Name, price, stock, and image path in records"],
        ["FR7", "Search and view own order", "OrderRepository checks customer ID"],
        ["FR8", "Combine text and image signals", "Weighted multimodal ranking"],
    ])
    heading(doc, "2.3 Non-functional requirements", 2)
    paragraph(doc, "Maintainability: the system is divided into presentation, application, and data modules, and services have narrow responsibilities. Reproducibility: the sample data and illustration generator are included, with exact run commands in README.md. Usability: command output reports the input, processing path, candidate count, and scores. Reliability: missing image files and empty voice transcripts produce explicit errors. Privacy: a sample order is returned only when both its ID and customer ID match; this is a prototype constraint, not a replacement for authentication. Performance: the in-memory design is appropriate only for a small product set.")

    # Page 4.
    new_page(doc, "3. Use Case Model")
    paragraph(doc, "The Customer may search for products using keyword, voice, or image input. These are three specialized forms of Search Product. The customer may then View Product. Separately, the customer may Search Order and View Order. The use cases state the customer's goals; the internal services that realize them appear in the component and sequence diagrams.")
    placeholder(doc, "VISUAL PARADIGM USE CASE DIAGRAM", "Insert your UML Use Case Diagram screenshot here. Include Customer; Search Product; Search by Keyword, Voice, and Image; Search Order; View Product; View Order.", 2.25)
    heading(doc, "3.1 Use case descriptions", 2)
    table(doc, ["Use case", "Precondition", "Main outcome"], [
        ["Search by Keyword", "Text entered", "Ranked product list"],
        ["Search by Voice", "Transcript supplied", "Transcript converted to text, then ranked products"],
        ["Search by Image", "Image file supplied", "Visually similar ranked products"],
        ["Search Order", "Order ID and customer ID supplied", "Matching order or not-found result"],
    ])
    paragraph(doc, "The voice branch uses a simulated transcript. In the UML model it may still be labeled voice search because the component boundary is the same as a future speech-to-text service. View Product and View Order are optional follow-up interactions after a search result is available.")

    # Page 5.
    new_page(doc, "4. Three-Layer Component Diagram with Packages")
    table(doc, ["UML package", "Components / modeled elements", "Responsibility"], [
        ["Presentation Layer", "SearchUI; VoiceInput; ImageUpload; SearchResultView", "Collect mode-specific input and present output"],
        ["Application / Intelligence Layer", "QueryService; SpeechService; ImageService; SearchService; RankingService", "Normalize, encode, retrieve, combine, and rank"],
        ["Data Layer", "ProductRepository; OrderRepository; VectorIndex; ProductDatabase; OrderDatabase; ImageStorage", "Provide product, order, and vector data"],
    ])
    placeholder(doc, "VISUAL PARADIGM THREE-LAYER COMPONENT DIAGRAM WITH PACKAGES", "Insert results/task3_component_diagram.png or a screenshot of Three-Layer Component Diagram from ass6.vpp here. It contains the three packages, components and their dependencies.", 2.20)
    heading(doc, "4.1 Components and dependencies", 2)
    paragraph(doc, "The Component Diagram shows the components recommended by the handout inside three packages. The Presentation Layer contains SearchUI, VoiceInput, ImageUpload, and SearchResultView. The Application / Intelligence Layer contains QueryService, SpeechService, ImageService, SearchService, and RankingService. The Data Layer contains ProductRepository, OrderRepository, VectorIndex, ProductDatabase, OrderDatabase, and ImageStorage. The storage components represent JSON files and image files in the prototype.")
    paragraph(doc, "Dependencies follow Presentation -> Application / Intelligence -> Data. SearchUI uses QueryService, SpeechService, and SearchService; SearchService uses the repositories, VectorIndex, ImageService, and RankingService. SearchUI does not read JSON files or the vector index directly. ImageService computes image vectors, while ProductRepository and OrderRepository read their JSON files. The Python entry point wires these components together.")
    heading(doc, "4.2 Retrieval versus ranking", 2)
    paragraph(doc, "Retrieval chooses a candidate set C(q) based on word overlap or cosine similarity. Ranking then orders C(q) using normalized relevance and a small stock signal. Keeping these steps separate makes the architecture visible in code and allows a later semantic retriever or business rule to be substituted without changing SearchUI.")

    # Page 6.
    new_page(doc, "5. Voice-Search Sequence Diagram")
    paragraph(doc, "This interaction demonstrates the required sequence with Customer as a UML actor and six system lifelines: SearchUI, SpeechService, QueryService, SearchService, ProductRepository, and RankingService. SearchUI has the boundary stereotype; SpeechService, QueryService, SearchService, and RankingService have the control stereotype; ProductRepository has the entity stereotype. The Customer supplies transcript text as the simulated audio input. SearchUI asks SpeechService to transcribe it, creates a voice query via QueryService, and asks SearchService to retrieve and rank products.")
    placeholder(doc, "VISUAL PARADIGM VOICE-SEARCH SEQUENCE DIAGRAM", "Insert results/task4_voice_search_sequence.png or a screenshot of Voice Search Sequence Diagram from ass6.vpp here. It shows calls and return messages for the transcript, query object, product candidates, ranked results, and displayed output.", 2.25)
    table(doc, ["Step", "Message / result"], [
        ["1–3", "Customer → SearchUI → SpeechService; recognized text returns"],
        ["4–5", "SearchUI → QueryService; normalized voice query returns"],
        ["6–8", "SearchUI → SearchService → ProductRepository; product list returns"],
        ["9–10", "SearchService → RankingService; ranked results return"],
        ["11–12", "SearchService → SearchUI → Customer; scored results displayed"],
    ])
    paragraph(doc, "The same application service receives a text query when the customer types a search. The image branch instead calls ImageService and VectorIndex. The voice sequence is therefore one concrete path through the shared architecture.")

    # Page 7.
    new_page(doc, "6. Python Implementation")
    heading(doc, "6.1 Modules and data", 2)
    table(doc, ["Module", "Purpose"], [
        ["presentation/search_ui.py", "Input methods and result display"],
        ["application/query_service.py", "Common query dictionaries"],
        ["application/speech_service.py", "Simulated speech-to-text"],
        ["application/image_service.py", "Pixel-based visual vector"],
        ["application/search_service.py", "Candidate retrieval and modality routing"],
        ["application/ranking_service.py", "Final score and ordering"],
        ["data/*_repository.py", "Read product and order JSON data"],
        ["data/vector_index.py", "In-memory cosine candidate search"],
    ])
    paragraph(doc, "The product dataset contains 12 items across shoes, bags, clothing, electronics, and accessories. Each record stores an ID, name, brand, category, color, price, stock, description, and image path. The sample images were generated specifically for this prototype. Two additional query illustrations introduce slight visual changes so that the image demonstration is not a byte-for-byte lookup of a stored file.")
    placeholder(doc, "OPTIONAL: SAMPLE PRODUCT AND QUERY IMAGES", "If you want images in the report, insert a small panel using files in data/images/ and data/query_images/. Identify the product image and the separate query illustration.", 1.25)
    heading(doc, "6.2 Execution flow", 2)
    paragraph(doc, "The main.py entry point builds repositories and vectors, then routes the selected command to SearchUI. For a text or voice query, SearchService normalizes words and retrieves products with overlapping tokens. An English price phrase such as ‘under 100’ filters products before ranking. For an image query, ImageService reads the file, generates a vector, and VectorIndex returns products with the greatest cosine similarity. SearchUI prints the final score and component scores.")

    # Page 8.
    new_page(doc, "7. Multimodal Search Method")
    heading(doc, "7.1 Common representation", 2)
    paragraph(doc, "QueryService produces a dictionary with a type field. Text and voice queries contain query text, image queries contain an image path, and combined queries contain both. This common envelope lets SearchService route the input without exposing the data layer to presentation details. Voice input first passes through SpeechService, which simply returns the provided transcript in this version.")
    heading(doc, "7.2 Text and image signals", 2)
    paragraph(doc, "Text retrieval tokenizes lowercase English words, removes a small stop-word list, performs minimal singular normalization, and checks the product's name, brand, color, category, and description. Matches in the name, brand, or color receive more weight. A small category bonus makes a shoe search favor shoes over unrelated black products. The resulting text relevance St is capped at 1.")
    paragraph(doc, "ImageService resizes each image to 96 × 96 pixels and extracts foreground color difference from white, a 12 × 12 spatial color grid, a 12 × 12 shape grid, mean foreground color, and foreground fraction. The descriptor is normalized to unit length. Cosine similarity is Si(q,p) = (zq · zp) / (||zq|| ||zp||). Because the vectors are normalized, VectorIndex can use their dot product to retrieve the closest products. This is a handcrafted image feature, not a learned image embedding.")
    heading(doc, "7.3 Candidate retrieval and ranking", 2)
    paragraph(doc, "For text and voice, candidate products have at least one matching query term. For image, VectorIndex supplies the ten nearest products. Multimodal search merges the two candidate lists by product ID. RankingService then uses the following scores, with stock normalized to min(stock/20, 1):")
    bullet(doc, "Text or voice: S = 0.98 St + 0.02 Sstock.")
    bullet(doc, "Image: S = 0.98 Si + 0.02 Sstock.")
    bullet(doc, "Text + image: S = 0.49 St + 0.49 Si + 0.02 Sstock.")
    paragraph(doc, "The combined weights sum to 1. The equal text and image weights follow the assignment's suggested λt = λi = 0.5, while reserving a small share for an observable business signal. In the demonstration, ‘black running shoes’ and a separate black-shoe illustration jointly rank Nike Black Running Shoes first with score 0.9911.")

    # Page 9.
    new_page(doc, "8. Experimental Results")
    heading(doc, "8.1 Evaluation setup", 2)
    paragraph(doc, "The evaluation uses 13 predefined queries. Success means the Top-1 product ID equals the expected ID; for the optional order lookup, the returned order ID must match. Query text, expected ID, actual result, score, and success are saved in results/evaluation.csv and results/evaluation.json. The demo transcript is in results/demo_results.txt. Image test queries use separate drawings from the stored product drawings.")
    rows = []
    for item in EVALUATION["results"]:
        label = item["query"]
        if len(label) > 31:
            label = label[:29] + "…"
        rows.append([
            label,
            item["mode"],
            str(item["expected_top1_id"]),
            str(item["actual_top1_id"] or "—"),
            str(item["ranking_score"] or "—"),
            "Yes" if item["success"] else "No",
        ])
    table(doc, ["Query", "Mode", "Expected", "Top-1", "Score", "Success"], rows)
    paragraph(doc, f"Successful queries: {EVALUATION['successful_queries']}/{EVALUATION['number_of_queries']}. Success rate = {EVALUATION['success_rate_percent']:.2f}%. The failed query ‘footwear for jogging’ returned no result even though Nike Black Running Shoes was the predefined target. Neither ‘footwear’ nor ‘jogging’ appears as a searchable term in the product record; the current text method does not interpret those synonyms.")
    # Page 10: Task 6 requires input, processing, products, and scores in the report itself.
    new_page(doc, "8.2 Demonstration of Required Search Modes")
    from main import build_services
    from presentation.search_ui import SearchUI

    ui = SearchUI(*build_services())
    demonstrations = [
        ("Text", "black running shoes", ui.search_text("black running shoes", limit=3), ""),
        ("Voice (simulated)", "find white sneakers",
         ui.search_voice("find white sneakers", limit=3), "simulated speech-to-text → "),
        ("Image", "data/query_images/query_red_backpack.png",
         ui.search_image(str(ROOT / "data/query_images/query_red_backpack.png"), limit=3), ""),
    ]
    demonstration_rows = []
    for mode, given_input, response, prefix in demonstrations:
        products = "\n".join(
            f"{rank}. {item['product']['name']} — {item['score']:.4f}"
            for rank, item in enumerate(response["results"], start=1)
        )
        demonstration_rows.append([mode, given_input, prefix + response["processing"], products])
    paragraph(doc, "Each row below gives the actual input, processing path, returned products, and final ranking scores from the executable prototype. The voice input is a transcript supplied to the simulated SpeechService. The image input is a separate query illustration, not a stored product file.")
    table(doc, ["Mode", "Input", "Processing", "Returned products and final scores"], demonstration_rows)
    paragraph(doc, "The full command-line transcript, including candidate counts and text/image/stock score components, is saved in results/demo_results.txt. The combined text-and-image demonstration is also present there and is explained in Section 7.")
    placeholder(doc, "PROGRAM SCREENSHOT: TEXT AND SIMULATED VOICE", "Insert a screenshot of `python main.py demo` showing the text and voice sections, including inputs and ranking scores.", 1.10)
    placeholder(doc, "PROGRAM SCREENSHOT: IMAGE SEARCH", "Insert a screenshot showing the image query path, returned products, image component scores, and final scores.", 1.10)

    # Page 11.
    new_page(doc, "9. Discussion and Limitations")
    heading(doc, "9.1 Interpretation", 2)
    paragraph(doc, "The results show that one application architecture can accept three input forms and return a consistently ranked product list. The separate query images ranked their intended products first, suggesting that the simple color-and-shape descriptor is adequate for the included illustrations. The experiment is small and controlled, so its 92.31% success rate should not be interpreted as expected performance on real customer photos or broad natural language.")
    placeholder(doc, "PROGRAM SCREENSHOT: MULTIMODAL SEARCH", "Insert a screenshot showing the combined text + image input and its ranked products with text, image, stock, and final scores.", 1.35)
    heading(doc, "9.2 Known limitations and possible extensions", 2)
    paragraph(doc, "Speech recognition is simulated, so the system does not handle microphone audio or recognition errors. The handcrafted image descriptor is sensitive to backgrounds and changes in viewpoint. Keyword search misses semantic equivalents such as ‘jogging’ for ‘running.’ The JSON repository and linear vector scan are appropriate only for a tiny dataset. Stock is a deliberately small ranking signal, and the price parser recognizes only simple English phrases. A future version could add a real speech-to-text model, learned image/text embeddings, a scalable vector index, authenticated order access, and a web interface.")
    paragraph(doc, "Order lookup in this classroom prototype checks the supplied customer ID but does not authenticate who supplied it. Production deployment would require account identity and authorization before returning order details. This matters because the Use Case Diagram includes order viewing, even though the principal assessment focus is product search.")

    # Page 12.
    new_page(doc, "10. Conclusion and Submission Map")
    paragraph(doc, "The project satisfies the main technical requirements through a working three-layer Python prototype. Text, simulated voice, and image input are mapped to one query format; retrieval selects candidates and ranking orders them. The implementation includes 12 products, real pixel-based image descriptors, a weighted multimodal mode, order lookup, and saved demonstration results. The evaluation records both successful examples and a concrete failure. The supplied ass6.vpp contains the Use Case Diagram, three-layer Component Diagram, and voice-search Sequence Diagram. The screenshots marked in yellow remain to be inserted.")
    heading(doc, "10.1 Requirements traceability", 2)
    table(doc, ["Assignment task", "Evidence in submission"], [
        ["Requirements analysis", "Section 2 and functional requirements table"],
        ["Use Case Diagram", "Existing diagram in ass6.vpp; Section 3 screenshot placeholder"],
        ["Three-layer component diagram", "Three-Layer Component Diagram in ass6.vpp; Section 4 screenshot placeholder"],
        ["Voice sequence diagram", "Voice Search Sequence Diagram in ass6.vpp; exported image in results/task4_voice_search_sequence.png; Section 5 screenshot placeholder"],
        ["Python implementation", "main.py, application/, presentation/, data/"],
        ["Three-mode demonstration", "Section 8.2 and results/demo_results.txt"],
        ["Experimental evaluation", "evaluate.py, results/evaluation.csv/json"],
    ])
    heading(doc, "10.2 Files to submit", 2)
    paragraph(doc, "Submit the exported PDF report, the completed Visual Paradigm project (ass6.vpp), Python source code, data/products.json and data/orders.json, sample PNG images, README.md, and the results/ folder. Keep the editable DOCX while inserting the figures and screenshots. The submission also contains VISUAL_PARADIGM_GUIDE.md as a diagram checklist.")
    heading(doc, "Reference", 2)
    paragraph(doc, "Course handout: Information System Analysis and Design, “Assignment 06: Multimodal Search System for E-Commerce,” 23-page PDF supplied with this assignment. All product data, sample illustrations, source code, and measured results in this report were created for this project.")

    doc.save(DESTINATION)
    print(f"Created {DESTINATION.name}")


if __name__ == "__main__":
    main()
